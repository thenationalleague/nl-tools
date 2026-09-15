"""Commercial Benchmarking — raw SurveyMonkey export -> the cleaned workbook.

build-benchmarks.py and build-benchmark-rows.py read a *cleaned* workbook:
sheet `Data`, one club per row, tidy headers, numbers as numbers. The seed's
docstring credits that cleaning to a clean_survey.py that never reached the
repo, so late returns had no route in. This is that step.

Column positions below are from the live SurveyMonkey export read on
15/09/2026 (100 responses, 94 columns, two header rows: the question, then
the sub-answer — "Response", "Other (please specify)", "Date / Time"…).
They are located by header text, not by index, so a re-export with the same
questions still maps; a changed question fails loudly.

What it does, per response:
  * keeps only the mapped answer columns — respondent IP, email, names and
    the club-representative block are never read, let alone written;
  * sector: the picked option, or the "Other (please specify)" text; "N/A" -> blank;
  * money and counts: "£12,000", "12k", "3,000+", "600 + VAT" -> a number;
    a range or prose ("300-500", "same as above") -> blank, and REPORTED;
  * contract length: "1 year", "Three Year", 2, "18 months" -> years;
    "rolling"/"ongoing" -> blank length + Rolling? = Yes;
  * commencement dates pass through (the parser reads Excel dates and
    DD/MM/YYYY strings).

Then per club, one row: responses with no answers at all are dropped (a club
that chose its name and stopped), and of the rest the fullest wins, newest
on a tie. --since keeps only responses ended on/after a date — the late
returns after a seed — and --clubs limits to named clubs.

Usage:
    python scripts/clean-survey-export.py <export.xlsx> <cleaned.xlsx> [--since YYYY-MM-DD] [--clubs "A,B"]

Nothing here is committed. Both files name clubs; the export also names
people. Keep them out of the repo and delete them when done.
"""
import sys, re, datetime, argparse
import openpyxl

CLEAN_HEADERS = ['Club', 'Division',
    'Top GA matchday ticket (£)', 'Top GA season ticket (£)',
    'Front Shirt — Sponsor Name', 'Front Shirt — Sector', 'Front Shirt — Income/season (£, ex-VAT)',
    'Front Shirt — Contract Length', 'Front Shirt — Commencement', 'Front Shirt — Rolling?',
    'Back Shirt — Sponsor Name', 'Back Shirt — Sector', 'Back Shirt — Income/season (£, ex-VAT)',
    'Back Shirt — Contract Length', 'Back Shirt — Commencement', 'Back Shirt — Rolling?',
    'Sleeve — Sponsor Name', 'Sleeve — Sector', 'Sleeve — Income/season (£, ex-VAT)',
    'Sleeve — Contract Length', 'Sleeve — Commencement', 'Sleeve — Rolling?',
] + [h % s for s in (1, 2, 3, 4) for h in (
    'Stand %d — Sponsor Name', 'Stand %d — Sector', 'Stand %d — Income/season (£, ex-VAT)',
    'Stand %d — Commencement', 'Stand %d — Contract Length')] + [
    'TV-facing board price/season (£)', 'Non-TV board price/season (£)',
    'Top matchday hospitality (£)', 'Top seasonal hospitality (£)',
    'Full-page programme advert/season (£)',
    'Total email database size', 'Opted-in to partner emails',
    'Programme format', 'Can email supporters?', 'Can email on behalf of partners?',
]

WORD_NUM = {'one': 1, 'two': 2, 'three': 3, 'four': 4, 'five': 5, 'six': 6, 'seven': 7, 'eight': 8, 'nine': 9, 'ten': 10,
            'a': 1, 'an': 1, 'single': 1}
NA = re.compile(r'^(n/?a|none|nil|-|–|—|tbc|tbd|\.|x)$', re.I)


def blank(v):
    return v is None or (isinstance(v, str) and (not v.strip() or NA.match(v.strip())))


def money(v):
    """Open-ended money/count answer -> number, or None. Second value says
    whether the text was unusable (so the caller can report it)."""
    if blank(v):
        return None, False
    if isinstance(v, bool):
        return None, True
    if isinstance(v, (int, float)):
        return float(v), False
    s = str(v).strip().lower().replace('£', '').replace(',', '')
    s = re.sub(r'\b(per|a|each)\s+(season|year|match|game)\b', '', s)
    s = re.sub(r'\+\s*vat\b|\bex\.?\s*vat\b|\binc\.?\s*vat\b|\bplus\s+vat\b|\bvat\b', '', s)
    s = s.replace('+', ' ').strip()
    m = re.fullmatch(r'(\d+(?:\.\d+)?)\s*k', s)
    if m:
        return float(m.group(1)) * 1000, False
    nums = re.findall(r'\d+(?:\.\d+)?', s)
    if len(nums) == 1 and re.fullmatch(r'[\d.\s]*', s):
        return float(nums[0]), False
    return None, True


def years(v):
    """Contract length -> (years or None, rolling bool, unusable bool)."""
    if blank(v):
        return None, False, False
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v), False, False
    s = str(v).strip().lower()
    if 'roll' in s or 'ongoing' in s or 'open' in s:
        return None, True, False
    for w, n in WORD_NUM.items():
        s = re.sub(r'\b%s\b' % w, str(n), s)
    m = re.search(r'(\d+(?:\.\d+)?)', s)
    if not m:
        return (1.0, False, False) if ('annual' in s or 'season' in s or 'year' in s) else (None, False, True)
    n = float(m.group(1))
    return (round(n / 12.0, 1) if 'month' in s else n), False, False


def sector(picked, other):
    if not blank(other):
        return str(other).strip()
    if blank(picked):
        return ''
    p = str(picked).strip()
    return '' if p.lower().startswith('other') else p


def find_cols(h1, h2):
    """Map each cleaned header to a raw column index by header text. Repeated
    questions (sector, commencement, length, income) are resolved by walking
    forward from the sponsor-name column that opens each block."""
    def idx(q, sub=None, start=0):
        for i in range(start, len(h1)):
            if h1[i] and str(h1[i]).strip().lower().startswith(q.lower()) and (sub is None or (h2[i] and str(h2[i]).strip().lower() == sub.lower())):
                return i
        raise KeyError('column not found: %r / %r' % (q, sub))
    def after(start, sub):
        for i in range(start, min(start + 8, len(h1))):
            if h2[i] and str(h2[i]).strip().lower() == sub.lower():
                return i
        raise KeyError('sub-column %r not found after %d' % (sub, start))
    C = {'Club': idx('Club name'), 'Division': idx('Division')}
    blocks = [('Front Shirt', 'Name of Front of Shirt Sponsor'), ('Back Shirt', 'Name of Back of Shirt Sponsor'), ('Sleeve', 'Name of Sleeve Sponsor')] + \
             [('Stand %d' % s, 'Name of Stand Sponsor - Stand %d' % s) for s in (1, 2, 3, 4)]
    for label, q in blocks:
        n = idx(q)
        C[label + ' — Sponsor Name'] = n
        C[label + ' — Sector'] = after(n, 'Response')
        C[label + ' — Sector (other)'] = after(n, 'Other (please specify)')
        C[label + ' — Commencement'] = after(n, 'Date / Time')
        # length then income are the two open-ended answers after the date
        d = C[label + ' — Commencement']
        C[label + ' — Contract Length'] = after(d + 1, 'Open-Ended Response')
        C[label + ' — Income/season (£, ex-VAT)'] = after(C[label + ' — Contract Length'] + 1, 'Open-Ended Response')
    C['TV-facing board price/season (£)'] = idx('How much do you charge for a seasonal TV facing')
    C['Non-TV board price/season (£)'] = idx('How much do you charge for a seasonal non TV')
    C['Programme?'] = idx('Do you produce a matchday programme')
    C['Programme format'] = idx('If yes, is the programme')
    C['Full-page programme advert/season (£)'] = idx('If yes, please give cost for a full page')
    C['Top GA matchday ticket (£)'] = idx('How much do you charge for your highest priced General Admission adult matchday')
    C['Top GA season ticket (£)'] = idx('How much do you charge for your highest priced General Admission adult season')
    C['Top matchday hospitality (£)'] = idx('How much do you charge for your highest priced matchday hospitality')
    C['Top seasonal hospitality (£)'] = idx('How much do you charge for your highest priced seasonal hospitality')
    C['Can email supporters?'] = idx('Do you currently have the ability to send email')
    C['Can email on behalf of partners?'] = idx('If yes, do you have the ability to send out email')
    C['Total email database size'] = idx('What is your current total email database size')
    C['Opted-in to partner emails'] = idx('Out of that total database')
    C['End Date'] = idx('End Date')
    C['Start Date'] = idx('Start Date')
    return C


def clean_row(r, C):
    """One raw response -> (dict keyed by CLEAN_HEADERS, [unusable (header, raw)])."""
    out, bad = {}, []
    out['Club'] = str(r[C['Club']]).strip() if r[C['Club']] else ''
    out['Division'] = str(r[C['Division']]).strip() if r[C['Division']] else ''
    def m(key):
        v, unusable = money(r[C[key]])
        if unusable:
            bad.append((key, r[C[key]]))
        out[key] = v
    for label in ('Front Shirt', 'Back Shirt', 'Sleeve'):
        nm = r[C[label + ' — Sponsor Name']]
        out[label + ' — Sponsor Name'] = '' if blank(nm) else str(nm).strip()
        out[label + ' — Sector'] = sector(r[C[label + ' — Sector']], r[C[label + ' — Sector (other)']])
        m(label + ' — Income/season (£, ex-VAT)')
        yrs, rolling, unusable = years(r[C[label + ' — Contract Length']])
        if unusable:
            bad.append((label + ' — Contract Length', r[C[label + ' — Contract Length']]))
        out[label + ' — Contract Length'] = yrs
        out[label + ' — Rolling?'] = 'Yes' if rolling else ('No' if yrs is not None else '')
        out[label + ' — Commencement'] = r[C[label + ' — Commencement']]
    for s in (1, 2, 3, 4):
        label = 'Stand %d' % s
        nm = r[C[label + ' — Sponsor Name']]
        out[label + ' — Sponsor Name'] = '' if blank(nm) else str(nm).strip()
        out[label + ' — Sector'] = sector(r[C[label + ' — Sector']], r[C[label + ' — Sector (other)']])
        m(label + ' — Income/season (£, ex-VAT)')
        yrs, rolling, unusable = years(r[C[label + ' — Contract Length']])
        if unusable:
            bad.append((label + ' — Contract Length', r[C[label + ' — Contract Length']]))
        out[label + ' — Contract Length'] = yrs
        out[label + ' — Commencement'] = r[C[label + ' — Commencement']]
    for key in ('TV-facing board price/season (£)', 'Non-TV board price/season (£)', 'Full-page programme advert/season (£)',
                'Top GA matchday ticket (£)', 'Top GA season ticket (£)', 'Top matchday hospitality (£)',
                'Top seasonal hospitality (£)', 'Total email database size', 'Opted-in to partner emails'):
        m(key)
    prog = r[C['Programme?']]
    fmt = r[C['Programme format']]
    out['Programme format'] = ('None' if (not blank(prog) and str(prog).strip().lower() == 'no') else
                               ('' if blank(fmt) else str(fmt).strip()))
    for key in ('Can email supporters?', 'Can email on behalf of partners?'):
        v = r[C[key]]
        out[key] = '' if blank(v) else str(v).strip()
    return out, bad


def answered(r, C):
    """Number of answer cells filled — anything from the first sponsor name on."""
    first = C['Front Shirt — Sponsor Name']
    return sum(1 for v in r[first:] if v not in (None, ''))


def load(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise SystemExit('export has no header rows')
    return rows[0], rows[1], rows[2:]


def clean(path, since=None, clubs=None):
    h1, h2, data = load(path)
    C = find_cols(h1, h2)
    kept, report = {}, []
    for r in data:
        club = r[C['Club']]
        if blank(club):
            continue
        club = str(club).strip()
        ended = r[C['End Date']]
        if since and (not isinstance(ended, datetime.datetime) or ended.date() < since):
            continue
        if clubs and club not in clubs:
            continue
        n = answered(r, C)
        if n == 0:
            report.append((club, ended, 0, 'empty response — club and division only, dropped'))
            continue
        cur = kept.get(club)
        if cur and (cur['n'] > n or (cur['n'] == n and cur['ended'] >= ended)):
            report.append((club, ended, n, 'superseded by a fuller or newer response'))
            continue
        if cur:
            report.append((club, cur['ended'], cur['n'], 'superseded by a fuller or newer response'))
        row, bad = clean_row(r, C)
        kept[club] = {'row': row, 'bad': bad, 'n': n, 'ended': ended}
    return kept, report


def write(kept, out):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    ws.append(CLEAN_HEADERS)
    for club in sorted(kept):
        ws.append([kept[club]['row'].get(h) for h in CLEAN_HEADERS])
    wb.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('export'); ap.add_argument('out')
    ap.add_argument('--since', help='keep responses ended on/after YYYY-MM-DD')
    ap.add_argument('--clubs', help='comma-separated club names to keep')
    a = ap.parse_args()
    since = datetime.date.fromisoformat(a.since) if a.since else None
    clubs = set(c.strip() for c in a.clubs.split(',')) if a.clubs else None
    kept, report = clean(a.export, since, clubs)
    for club, ended, n, why in report:
        print('SKIP %-28s %s  %2d answers — %s' % (club, str(ended)[:10], n, why))
    for club in sorted(kept):
        k = kept[club]
        print('KEEP %-28s %s  %2d answers%s' % (club, str(k['ended'])[:10], k['n'],
              ('  — %d cell%s need a look' % (len(k['bad']), '' if len(k['bad']) == 1 else 's')) if k['bad'] else ''))
        for key, raw in k['bad']:
            print('       could not read %-40s %r  (left blank; fix in the editor after import)' % (key, raw))
    write(kept, a.out)
    print('\nWrote %s: %d club%s' % (a.out, len(kept), '' if len(kept) == 1 else 's'))
    print('Next: python scripts/build-benchmark-rows.py %s' % a.out)


if __name__ == '__main__':
    main()
